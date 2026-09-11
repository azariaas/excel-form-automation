"""Leitura, marcação e escrita da planilha de entrada."""

from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)


def map_headers(ws: Worksheet, header_row: int) -> dict[str, int]:
    """Mapeia o texto de cada cabeçalho (normalizado) para o número da coluna."""
    headers = {}
    for cell in ws[header_row]:
        if cell.value:
            headers[str(cell.value).strip().lower()] = cell.column
    return headers


def is_row_marked(cell, highlight_color: str) -> bool:
    """Verifica se a célula já está destacada com a cor de 'processado'."""
    fill = cell.fill
    if not fill or not fill.fgColor:
        return False
    color = fill.fgColor.rgb if fill.fgColor.type == "rgb" else ""
    if not color:
        return False
    return color.upper().endswith(highlight_color.upper())


def mark_row(row, highlight_color: str) -> None:
    """Pinta todas as células da linha com a cor de 'processado'."""
    fill = PatternFill("solid", fgColor=highlight_color)
    for cell in row:
        cell.fill = fill


def get_row_data(ws: Worksheet, row_number: int, headers: dict[str, int], fields: list[str]) -> dict[str, str]:
    """Extrai os valores de uma linha para os campos configurados."""
    data = {}
    for field in fields:
        column = headers.get(field.lower())
        if column is None:
            data[field] = ""
            continue
        value = ws.cell(row=row_number, column=column).value
        data[field] = str(value).strip() if value is not None else ""
    return data


def load_workbook_file(path: Path):
    if not path.exists():
        raise FileNotFoundError(
            f"Planilha não encontrada: {path}. Gere uma planilha de exemplo com "
            "'python scripts/generate_sample_data.py' ou ajuste 'spreadsheet_path' na configuração."
        )
    return load_workbook(path)


def save_workbook_file(wb, path: Path) -> None:
    wb.save(path)
